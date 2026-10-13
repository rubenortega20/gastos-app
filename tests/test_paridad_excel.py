"""Paridad con el Excel original. Se omite si no defines XLSX_ORIGINAL=ruta/Orden_def.xlsx.
Valida que la app reproduzca los números cacheados del Dashboard/Presupuesto/Cuotas de oct-2026."""
import os
import unittest

import openpyxl

from app import calc, db, importer

RUTA = os.environ.get("XLSX_ORIGINAL")


@unittest.skipUnless(RUTA and os.path.exists(RUTA), "define XLSX_ORIGINAL para correr la paridad")
class TestParidad(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.c = db.conectar(":memory:")
        importer.importar(cls.c, RUTA)
        cls.v = openpyxl.load_workbook(RUTA, data_only=True)
        cls.p = cls.v["Dashboard"]["B2"].value.strftime("%Y-%m")
        cls.d = calc.dashboard(cls.c, cls.p)

    def test_kpis_dashboard(self):
        ws = self.v["Dashboard"]
        esperado = {"ingreso": "A5", "ahorro_objetivo": "D5", "ahorro_real_excel": "A8", "disponible": "D8",
                    "tasa_ahorro": "B10", "meta_ahorro": "E10", "cuotas_mensuales": "B11", "saldo_deudas": "E11",
                    "total_disponible": "D27", "saldo_cuenta": "D30"}
        for k, celda in esperado.items():
            with self.subTest(k):
                self.assertAlmostEqual(self.d[k], ws[celda].value, places=4)

    def test_presupuesto_por_categoria(self):
        ws = self.v["Presupuesto"]
        for i, fila in enumerate(self.d["categorias"]):
            r = 4 + i
            with self.subTest(fila["categoria"]):
                self.assertEqual(fila["categoria"], ws[f"A{r}"].value)
                for k, col in (("presupuesto", "B"), ("fijo", "C"), ("movimientos", "D"), ("real", "F"), ("disponible", "G")):
                    self.assertAlmostEqual(fila[k], ws[f"{col}{r}"].value, places=4)

    def test_facturacion_tarjetas(self):
        ws = self.v["Dashboard"]
        esperado = {ws[f"A{r}"].value: ws[f"B{r}"].value for r in (29, 30, 31)}
        obtenido = {t["tarjeta"]: t["total"] for t in self.d["tarjetas"]}
        for nombre, valor in esperado.items():
            with self.subTest(nombre):
                self.assertAlmostEqual(obtenido[nombre], valor, places=4)

    def test_cuotas_todos_los_meses(self):
        ws = self.v["Cuotas"]
        for r in range(4, 15):
            per = ws[f"J{r}"].value.strftime("%Y-%m")
            cs = calc.cuotas_mes(self.c, per)
            with self.subTest(per):
                self.assertAlmostEqual(sum(x["cuota_mensual"] for x in cs), ws[f"K{r}"].value)
                self.assertAlmostEqual(sum(x["saldo_pendiente"] for x in cs), ws[f"L{r}"].value)


if __name__ == "__main__":
    unittest.main()
